import os
import json
import logging
import re
import subprocess
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Dict, Optional, List, Tuple, Any
from dataclasses import dataclass, asdict
import base64

# --- Dependency Imports with Graceful Fallbacks ---
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import ollama
except ImportError:
    ollama = None

try:
    import markdown
except ImportError:
    markdown = None

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

import io
try:
    from PIL import Image
except ImportError:
    Image = None


logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

@dataclass
class DocumentSummary:
    document_name: str
    document_type: str
    important_details: Dict[str, Any] = None
    error: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        data = {
            "document_name": self.document_name,
            "document_type": self.document_type,
            "important_details": self.important_details or {},
        }
        if self.error:
            data["error"] = self.error
        return data


class DocumentIntelligenceEngine:
    """Classifies the document and extracts relevant procurement/financial fields (V5 Style)."""

    def __init__(self, model: str = "gemma3:4b") -> None:
        self.model = model

    def analyze(self, document_name: str, plain_text: str) -> DocumentSummary:
        if not ollama:
            return DocumentSummary(
                document_name=document_name,
                document_type="Generic Supporting Document",
                error="'ollama' library is not installed.",
            )

        lower_text = plain_text[:1000].lower()
        lower_name = document_name.lower()
        is_complex = any(k in lower_text or k in lower_name for k in [
            "financial statement", "technical spec", "tender", "balance sheet", 
            "profit and loss", "specification", "schedule of requirement"
        ])
        limit = 2500 if is_complex else 1500
        truncated_text = plain_text[:limit]

        prompt = f'''

Return ONLY valid raw JSON.

No markdown.

No explanation.

No thinking.

No commentary.

No ```json wrapper.

Do NOT guess.

Do NOT hallucinate.

Do NOT invent values.

Do NOT force classification.

If information is unclear, omit the field.

If document does not strongly match a known category,
return:

Generic Supporting Document

Only include fields that are clearly visible in the document.

Never generate placeholder values like:

N/A
Unknown
Not Available
Buyer
Supplier
M/s ........
__________

Do NOT return dotted placeholders or blank template text.

If a field contains only blanks, dots, placeholders, or template text,
IGNORE that field completely.

==================================================
DOCUMENT TYPES
==================================================

Allowed document types:

Aadhaar Card
PAN Card
GST Certificate
Bank Guarantee
Financial Statement
Invoice
Purchase Order
Tender Document
Technical Specifications
Compliance Document
Bank Document
Generic Supporting Document

==================================================
STRICT CLASSIFICATION RULES
==================================================

Only classify as Aadhaar Card if:
- Aadhaar heading exists
- OR UIDAI style layout exists
- OR a valid 12-digit Aadhaar number exists

Only classify as PAN Card if:
- PAN heading exists
- OR valid PAN format exists

Only classify as GST Certificate if:
- GSTIN exists
- OR GST registration format exists

Only classify as Purchase Order if:
- PO Number
- Contract Number
- Supply Order
- Procurement Order

clearly exists

Only classify as Technical Specifications if:
- item code
- part number
- material composition
- dimensions
- technical values
- specification table
- grade
- composition

clearly exists

Only classify as Tender Document if:
- bid
- tender clause
- scope of work
- annexure
- bidder eligibility

clearly exists

Otherwise:

Generic Supporting Document

==================================================
CRITICAL EXTRACTION RULES
==================================================

1. Aadhaar Card

MANDATORY:
Extract BOTH:

- full_person_name
- aadhaar_number

aadhaar_number MUST be a valid 12-digit number

Accepted formats:

XXXX XXXX XXXX
XXXXXXXXXXXX

If Aadhaar number exists,
DO NOT miss it.

--------------------------------------------------

AADHAAR NAME RULE (VERY IMPORTANT)

Do NOT pick random single words like:

"मधुर"

or partial OCR fragments.

Do NOT extract only one Hindi word if the full name exists.

You must identify the ACTUAL PERSON NAME.

Use semantic understanding.

Usually Aadhaar contains:

1. English full name
2. Hindi/local language name
3. Father/Mother name
4. UIDAI text

Extract the REAL primary person name,
preferably the English full name if present.

Priority:

English full name
>
Full Hindi name
>
Readable full native-language name

Never prefer a single isolated word over a full proper name.

Bad example:

"name": "मधुर"

Good example:

"name": "Dayanand Ningayya Mayur"

or

full proper readable full name

Only extract complete personal name.

Not fragments.

--------------------------------------------------

2. PAN Card

Extract:

- name
- pan_number

pan_number must match PAN format

--------------------------------------------------

3. GST Certificate

Extract:

- legal_name
- gstin / registration_number

--------------------------------------------------

4. Purchase Order

Extract only if clearly visible:

- contract_number
- po_number
- supplier_name

Do NOT extract fake supplier address
from dotted blank templates.

Ignore placeholder addresses.

--------------------------------------------------

5. Technical Specifications

Extract only real technical values like:

- material_composition
- grade
- dimensions
- specification_details
- part_number
- item_code

==================================================
VERY IMPORTANT ADDRESS RULE
==================================================

Never extract addresses like:

"M/s ............."

or dotted blank forms

or template placeholders.

If supplier address is incomplete template text,
DO NOT include supplier_address.

Only include address if real readable address exists.

==================================================
DOCUMENT CONTENT
==================================================

{truncated_text}

==================================================
OUTPUT FORMAT
==================================================

{{
  "document_name": "{document_name}",
  "document_type": "document type",
  "important_details": {{
    "field": "value"
  }}
}}

Return ONLY JSON.
'''

        try:
            response = ollama.generate(
                model=self.model,
                prompt=prompt,
                stream=False,
                options={
                    "num_ctx": 2048,
                    "num_predict": 512,
                    "temperature": 0,
                    "top_p": 1
                },
            )
            raw = response.get("response", "").strip()
            
            logger.info("RAW MODEL RESPONSE:\n%.500s", raw)
            
            return self._parse_response(document_name, raw)

        except Exception as exc:
            logger.error("Intelligence engine error for '%s': %s", document_name, exc)
            return DocumentSummary(
                document_name=document_name,
                document_type="Generic Supporting Document",
                error=str(exc),
            )

    def stop_ocr_model(self) -> None:
        """Stop the glm-ocr model in Ollama to reclaim system memory."""
        try:
            ollama.generate(model="glm-ocr:q8_0", keep_alive=0)
            logger.info("OCR model explicitly unloaded from VRAM.")
        except Exception as exc:
            logger.debug("Could not stop OCR model: %s", exc)

    @staticmethod
    def _clean_raw_response(raw: str) -> str:
        # 1. Strip think blocks
        cleaned = re.sub(r"<think>.*?</think>", "", raw, flags=re.DOTALL | re.I).strip()
        # 2. Strip json wrappers
        cleaned = re.sub(r"^```[a-zA-Z]*\s*", "", cleaned, flags=re.M)
        cleaned = re.sub(r"```", "", cleaned).strip()
        # 3. Explanation before first {
        brace_pos = cleaned.find("{")
        if brace_pos > 0:
            cleaned = cleaned[brace_pos:]
        # 4. Text after last }
        last_brace = cleaned.rfind("}")
        if last_brace != -1:
            cleaned = cleaned[:last_brace+1]
        # 5. Empty lines
        cleaned = os.linesep.join([s for s in cleaned.splitlines() if s.strip()])
        return cleaned.strip()

    @classmethod
    def _parse_response(cls, document_name: str, raw: str) -> DocumentSummary:
        cleaned = cls._clean_raw_response(raw)

        try:
            data = json.loads(cleaned)
            return DocumentSummary(
                document_name=data.get("document_name", document_name),
                document_type=data.get("document_type", "Generic Supporting Document"),
                important_details=data.get("important_details", {}),
            )
        except json.JSONDecodeError:
            pass

        json_match = re.search(r"\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", cleaned, re.DOTALL)
        if json_match:
            try:
                data = json.loads(json_match.group(0))
                return DocumentSummary(
                    document_name=data.get("document_name", document_name),
                    document_type=data.get("document_type", "Generic Supporting Document"),
                    important_details=data.get("important_details", {}),
                )
            except json.JSONDecodeError:
                pass

        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start != -1 and end != -1 and end > start:
            try:
                data = json.loads(cleaned[start : end + 1])
                return DocumentSummary(
                    document_name=data.get("document_name", document_name),
                    document_type=data.get("document_type", "Generic Supporting Document"),
                    important_details=data.get("important_details", {}),
                )
            except json.JSONDecodeError:
                pass

        logger.warning("Could not parse intelligence response for '%s'. Raw: %.300s", document_name, raw)
        return DocumentSummary(
            document_name=document_name,
            document_type="Generic Supporting Document",
            error=f"JSON parse failed. Raw snippet: {raw[:200]}",
        )


# =============================================================
# Pipeline Orchestration (V7 Agent)
# =============================================================

